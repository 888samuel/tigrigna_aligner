from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class Phrase:
    index: int
    text: str
    normalized: str


_PUNCT_RE = re.compile(
    r"[\s"
    r"\.,;:!?\-—–…\"'“”‘’\(\)\[\]\{\}<>/"
    r"፡።፣፤፥፦፧፨"  # Ethiopic punctuation
    r"]+",
    re.UNICODE,
)

_SENTENCE_SPLIT_RE = re.compile(
    r"(?<=[።!?\.؟!])\s+|(?<=[:፡፣፤])\s+",
    re.UNICODE,
)

TEXT_EXTENSIONS = {
    ".txt",
    ".text",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".srt",
    ".vtt",
}
DOCX_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}
JSON_EXTENSIONS = {".json"}
SUPPORTED_SCRIPT_EXTENSIONS = (
    TEXT_EXTENSIONS | DOCX_EXTENSIONS | PDF_EXTENSIONS | JSON_EXTENSIONS | {".rtf"}
)


def normalize_tigrinya(text: str) -> str:
    """NFC normalize and strip punctuation/extra space for matching."""
    text = unicodedata.normalize("NFC", text.strip())
    text = _PUNCT_RE.sub(" ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text.casefold()


def load_script(path: str | Path) -> list[Phrase]:
    """Load phrases from txt/md/csv/srt/vtt/json/docx/pdf/rtf."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Script not found: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_SCRIPT_EXTENSIONS:
        # Try plain-text decode for unknown extensions (e.g. .tig, no extension)
        try:
            raw = _read_plain_text(path)
        except UnicodeDecodeError as exc:
            supported = ", ".join(sorted(SUPPORTED_SCRIPT_EXTENSIONS))
            raise ValueError(
                f"Unsupported script format '{ext or '(none)'}'. "
                f"Supported: {supported}. Or save as UTF-8 text."
            ) from exc
    elif ext in DOCX_EXTENSIONS:
        raw = _extract_docx(path)
    elif ext in PDF_EXTENSIONS:
        raw = _extract_pdf(path)
    elif ext in JSON_EXTENSIONS:
        raw = _extract_json(path)
    elif ext == ".rtf":
        raw = _extract_rtf(path)
    elif ext in {".srt", ".vtt"}:
        raw = _extract_subtitle(path)
    else:
        raw = _read_plain_text(path)

    phrases = phrases_from_text(raw)
    if not phrases:
        raise ValueError(f"No phrases found in script: {path}")
    return phrases


def phrases_from_text(raw: str) -> list[Phrase]:
    """Turn extracted document text into one Phrase per spoken line/paragraph."""
    raw = unicodedata.normalize("NFC", raw.replace("\r\n", "\n").replace("\r", "\n"))
    lines = _candidate_lines(raw)
    phrases: list[Phrase] = []
    for line in lines:
        text = line.strip()
        if not text:
            continue
        if text.startswith("#") or text.startswith("//"):
            continue
        # Skip PDF page markers / common junk
        if re.fullmatch(r"(?i)page\s+\d+(\s+of\s+\d+)?", text):
            continue
        phrases.append(
            Phrase(
                index=len(phrases),
                text=text,
                normalized=normalize_tigrinya(text),
            )
        )
    return phrases


def _candidate_lines(raw: str) -> list[str]:
    """Prefer explicit newlines; otherwise split long blocks on sentence punctuation."""
    rough = [ln.strip() for ln in raw.splitlines()]
    non_empty = [ln for ln in rough if ln]
    if not non_empty:
        return []

    # If the file already looks phrase-per-line, keep it.
    if len(non_empty) >= 2:
        return non_empty

    # Single blob (common for PDF/DOCX paste): split into sentences/phrases.
    blob = non_empty[0]
    parts = [p.strip() for p in _SENTENCE_SPLIT_RE.split(blob) if p.strip()]
    return parts if parts else non_empty


def _read_plain_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "DOCX support requires python-docx. Run: pip install python-docx"
        ) from exc

    doc = Document(str(path))
    blocks: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append(text)
    # Also pull table cells (sometimes scripts live in tables)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    blocks.append(text)
    return "\n".join(blocks)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError(
            "PDF support requires pypdf. Run: pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    blocks: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            blocks.append(text)
    return "\n".join(blocks)


def _extract_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(data, list):
        lines: list[str] = []
        for item in data:
            if isinstance(item, str):
                lines.append(item)
            elif isinstance(item, dict):
                text = item.get("text") or item.get("phrase") or item.get("line")
                if text:
                    lines.append(str(text))
        return "\n".join(lines)
    if isinstance(data, dict):
        for key in ("phrases", "lines", "segments", "script"):
            value = data.get(key)
            if isinstance(value, list):
                return _extract_json_listish(value)
            if isinstance(value, str):
                return value
        if "text" in data:
            return str(data["text"])
    raise ValueError(
        "JSON script must be a string list, or an object with "
        "phrases/lines/segments/script/text"
    )


def _extract_json_listish(value: list) -> str:
    lines: list[str] = []
    for item in value:
        if isinstance(item, str):
            lines.append(item)
        elif isinstance(item, dict):
            text = item.get("text") or item.get("phrase") or item.get("line")
            if text:
                lines.append(str(text))
    return "\n".join(lines)


def _extract_rtf(path: Path) -> str:
    raw = _read_plain_text(path)
    try:
        from striprtf.striprtf import rtf_to_text

        return rtf_to_text(raw)
    except ImportError:
        # Lightweight fallback: strip common RTF control words
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
        text = re.sub(r"\\[a-zA-Z]+\d* ?", " ", text)
        text = text.replace("{", " ").replace("}", " ")
        return re.sub(r"\s+", " ", text).strip()


def _extract_subtitle(path: Path) -> str:
    raw = _read_plain_text(path)
    lines: list[str] = []
    for line in raw.splitlines():
        text = line.strip()
        if not text:
            continue
        if text.upper().startswith("WEBVTT"):
            continue
        if re.fullmatch(r"\d+", text):
            continue
        if re.search(r"\d{2}:\d{2}:\d{2}[.,]\d{3}\s*-->", text):
            continue
        if text.startswith("NOTE"):
            continue
        lines.append(text)
    return "\n".join(lines)
